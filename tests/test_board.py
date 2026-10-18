import re

import pytest

from draftboard import load_config
from draftboard.board import build_workbook
from draftboard.valuation import display_stats, value_all

FUNCS = {
    "AND", "ARRAY_CONSTRAIN", "ROWS", "SEQUENCE", "COUNTIFS", "MIN", "OR", "TRIM", "CHOOSE", "COUNTIF", "FALSE", "FILTER", "IF", "IFERROR", "INDEX",
    "ISNUMBER", "ISODD", "LARGE", "LEFT", "MATCH", "MAX", "MEDIAN", "MINIFS", "MOD", "NORM",
    "DIST", "NOT", "RANK", "ROUNDUP", "SEARCH", "SORT", "SUM", "SUMPRODUCT", "TEXTJOIN", "TRUE",
    "UNIQUE", "Log", "Players", "Config", "Board",
}


@pytest.fixture
def workbook(synthetic_players):
    cfg = load_config()
    lg = cfg["league"]
    vals = value_all(synthetic_players, lg["team_count_options"], lg["rounds"], cfg["valuation"])
    vals = vals.head(240).reset_index(drop=True)
    return build_workbook(vals, display_stats(vals), cfg), cfg


def _formulas(wb):
    for s in wb.sheets:
        for (r, c), v in s.cells.items():
            if isinstance(v, str) and v.startswith("="):
                yield s.name, r, c, v


def test_every_identifier_is_a_function_or_named_range(workbook):
    wb, _ = workbook
    known = FUNCS | set(wb.named)
    unknown = set()
    for _, _, _, f in _formulas(wb):
        body = re.sub(r'"[^"]*"', "", f)                      # drop string literals
        body = re.sub(r"'[^']*'!", "", body)
        for tok in re.findall(r"[A-Za-z_][A-Za-z_0-9.]*", body):
            if re.fullmatch(r"\$?[A-Z]{1,3}\$?\d+", tok) or re.fullmatch(r"[A-Z]{1,3}", tok):
                continue                                      # cell refs / column letters
            for part in tok.split("."):
                if part and part not in known:
                    unknown.add(part)
    assert not unknown, unknown


def test_no_hardcoded_team_count(workbook):
    wb, _ = workbook
    for name, r, c, f in _formulas(wb):
        # only Log row bounds (max picks = 196) may contain 196/197; 14 only as a CHOOSE index
        assert not re.search(r"(?<![\d$A-Z])(168|182|196)(?!\d)", f), (name, r, c, f)
        assert "/14" not in f and "*14" not in f, (name, r, c, f)


def test_log_has_max_picks_and_snake(workbook):
    wb, cfg = workbook
    log = wb.sheet("Log")
    assert log.cells[(196, 0)] == 196
    assert (197, 0) not in log.cells
    assert log.validations[0].source.startswith("=Players!$A$2:$A$")


def test_xlsx_renders(workbook, tmp_path):
    from openpyxl import load_workbook

    from draftboard.xlsx import render
    wb, _ = workbook
    out = render(wb, tmp_path / "b.xlsx")
    x = load_workbook(out)
    assert x.sheetnames == ["Config", "Board", "Log", "Players", "My Team", "Cheatsheet"]
    assert "NumTeams" in x.defined_names
    assert x["Log"]["C2"].value.startswith("=IF(ISODD(")


def snake_slot(pick, n):
    rnd = -(-pick // n)
    return (pick - 1) % n + 1 if rnd % 2 else n - (pick - 1) % n


@pytest.mark.parametrize("n,turns", [(12, [(12, 12), (13, 12), (24, 1), (25, 1)]),
                                     (13, [(13, 13), (14, 13), (26, 1), (27, 1)]),
                                     (14, [(14, 14), (15, 14), (28, 1), (29, 1)])])
def test_snake_formula_reference(n, turns):
    for pick, slot in turns:
        assert snake_slot(pick, n) == slot
