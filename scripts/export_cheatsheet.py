"""Part 4.4: render the Cheatsheet tab to out/cheatsheet.pdf (via LibreOffice, no API needed)."""
from __future__ import annotations

import argparse
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from draftboard import ROOT, load_config  # noqa: E402
from draftboard.model import Workbook  # noqa: E402
from draftboard.xlsx import render  # noqa: E402
from scripts.build_sheet import load_workbook_model  # noqa: E402


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--config", default=None)
    args = ap.parse_args()
    cfg = load_config(args.config)
    full = load_workbook_model(cfg)
    wb = Workbook(sheets=[full.sheet("Cheatsheet")])
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        sys.exit("LibreOffice not found; print the Cheatsheet tab from Google Sheets instead "
                 "(File > Download > PDF, 'Current sheet', landscape, fit to width).")
    out_dir = ROOT / "out"
    with tempfile.TemporaryDirectory() as tmp:
        x = render(wb, Path(tmp) / "cheatsheet.xlsx")
        from openpyxl import load_workbook
        xb = load_workbook(x)
        ws = xb.active
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:4"
        ws.print_options.gridLines = True
        xb.save(x)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out_dir), str(x)],
                       check=True, capture_output=True)
    print(f"wrote {out_dir / 'cheatsheet.pdf'}")


if __name__ == "__main__":
    main()
