"""Render a Workbook model to .xlsx (fallback: upload to Drive, open as Google Sheet).

Formulas are written as Google Sheets text. Excel won't evaluate FILTER/SORT/
ARRAY_CONSTRAIN as written, but Google Sheets does after import.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook as XWorkbook
from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .model import Workbook, parse_range


def _hex(c: str) -> str:
    return "FF" + c.lstrip("#").upper()


def render(wb: Workbook, path: Path) -> Path:
    xb = XWorkbook()
    xb.remove(xb.active)
    for s in wb.sheets:
        ws = xb.create_sheet(s.name)
        for (r, c), v in s.cells.items():
            if v == "":
                continue
            ws.cell(row=r + 1, column=c + 1, value=v)
        for col, px in s.widths.items():
            ws.column_dimensions[get_column_letter(col + 1)].width = max(3, px / 7)
        for col in s.hidden_cols:
            ws.column_dimensions[get_column_letter(col + 1)].hidden = True
        if s.frozen_rows or s.frozen_cols:
            ws.freeze_panes = ws.cell(row=s.frozen_rows + 1, column=s.frozen_cols + 1)
        if s.tab_color:
            ws.sheet_properties.tabColor = s.tab_color.lstrip("#").upper()
        for rng, f in s.formats:
            r0, c0, r1, c1 = parse_range(rng)
            for row in ws.iter_rows(min_row=r0 + 1, max_row=r1 + 1, min_col=c0 + 1, max_col=c1 + 1):
                for cell in row:
                    if f.number:
                        cell.number_format = f.number
                    if f.bold or f.italic or f.size or f.fg:
                        cell.font = Font(bold=f.bold or cell.font.bold, italic=f.italic,
                                         size=f.size or cell.font.size,
                                         color=_hex(f.fg) if f.fg else cell.font.color)
                    if f.bg:
                        cell.fill = PatternFill("solid", fgColor=_hex(f.bg))
                    if f.align or f.wrap:
                        cell.alignment = Alignment(horizontal=(f.align or "general").lower(),
                                                   wrap_text=f.wrap)
        for rng in s.merges:
            ws.merge_cells(rng)
        for v in s.validations:
            if v.kind == "checkbox":
                dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
            elif v.kind == "list":
                dv = DataValidation(type="list", formula1='"' + ",".join(map(str, v.source)) + '"')
            else:
                dv = DataValidation(type="list", formula1=v.source.lstrip("="))
            dv.showErrorMessage = v.strict
            dv.errorStyle = "stop"
            ws.add_data_validation(dv)
            dv.add(v.rng)
        for cf in s.cond_formats:
            if cf.kind == "scale":
                rule = ColorScaleRule(start_type="min", start_color=_hex(cf.min_color),
                                      mid_type="num", mid_value=cf.mid_value, mid_color=_hex(cf.mid_color),
                                      end_type="max", end_color=_hex(cf.max_color))
            else:
                font = Font(color=_hex(cf.fg) if cf.fg else None, bold=cf.bold or None,
                            strike=cf.strike or None)
                fill = PatternFill("solid", bgColor=_hex(cf.bg)) if cf.bg else None
                rule = FormulaRule(formula=[cf.formula.lstrip("=")], font=font, fill=fill)
            ws.conditional_formatting.add(cf.rng, rule)
    for name, (sheet, rng) in wb.named.items():
        xb.defined_names[name] = DefinedName(name, attr_text=f"'{sheet}'!{rng}")
    path.parent.mkdir(parents=True, exist_ok=True)
    xb.save(path)
    return path
